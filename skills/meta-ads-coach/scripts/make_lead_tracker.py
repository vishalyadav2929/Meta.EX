#!/usr/bin/env python3
"""Create a lead-tracker Excel sheet with a per-source summary.
Usage: python3 make_lead_tracker.py lead-tracker.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

out = sys.argv[1] if len(sys.argv) > 1 else "lead-tracker.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Leads"
cols = ["Date", "Name", "Phone", "Event date", "City", "Source (ad / keyword)", "Status", "Booking ₹", "Notes"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for col, w in zip("ABCDEFGHI", [12, 20, 15, 12, 14, 26, 12, 12, 30]):
    ws.column_dimensions[col].width = w
dv = DataValidation(type="list", formula1='"New,Talking,Booked,Lost"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("G2:G1000")
ws.freeze_panes = "A2"

s = wb.create_sheet("Summary")
s.append(["Source", "Leads", "Booked", "Booking ₹", "Ad spend ₹", "Cost per lead", "Cost per booking"])
for c in s[1]:
    c.font = Font(bold=True)
for r in range(2, 22):
    s[f"B{r}"] = f'=IF(A{r}="","",COUNTIF(Leads!F:F,A{r}))'
    s[f"C{r}"] = f'=IF(A{r}="","",COUNTIFS(Leads!F:F,A{r},Leads!G:G,"Booked"))'
    s[f"D{r}"] = f'=IF(A{r}="","",SUMIFS(Leads!H:H,Leads!F:F,A{r}))'
    s[f"F{r}"] = f'=IF(OR(A{r}="",E{r}="",B{r}=0),"",E{r}/B{r})'
    s[f"G{r}"] = f'=IF(OR(A{r}="",E{r}="",C{r}=0),"",E{r}/C{r})'
s["A2"], s["A3"], s["A4"] = "Baraat ad", "Sangeet ad", "Organic / referral"
s.column_dimensions["A"].width = 26
wb.save(out)
print(f"Saved {out}")
